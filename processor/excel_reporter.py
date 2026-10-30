"""
Excel Reporter Module
Generates an executive-styled multi-tab Excel workbook with KPI summaries and formatted currency cells.
"""

from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def generate_styled_excel_report(
    df: pd.DataFrame, metrics: Dict[str, Any], output_path: str
) -> None:
    """Creates a professional 2-sheet executive Excel workbook."""
    wb = openpyxl.Workbook()

    # Style definitions
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    kpi_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    currency_format = '#,##0.00 "zł"'
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # ==========================================
    # SHEET 1: Executive KPI Summary
    # ==========================================
    ws_kpi = wb.active
    ws_kpi.title = "Executive Summary"
    ws_kpi.views.sheetView[0].showGridLines = True

    # Title
    ws_kpi["A1"] = "ERP Sales & Invoicing Executive Summary"
    ws_kpi["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")

    kpis = [
        ("Total Netto Revenue", metrics["total_revenue_netto"], True),
        ("Total VAT Collected", metrics["total_vat"], True),
        ("Total Brutto Turnover", metrics["total_revenue_brutto"], True),
        ("Clean Invoices Processed", metrics["clean_rows"], False),
        ("Duplicates Removed", metrics["duplicates_removed"], False),
        ("Invalid Tax IDs (NIP) Detected", metrics["invalid_nips"], False),
    ]

    ws_kpi["A3"] = "Key Performance Indicator (KPI)"
    ws_kpi["B3"] = "Value"
    ws_kpi["A3"].fill = header_fill
    ws_kpi["B3"].fill = header_fill
    ws_kpi["A3"].font = header_font
    ws_kpi["B3"].font = header_font

    row_idx = 4
    for label, val, is_currency in kpis:
        c1 = ws_kpi.cell(row=row_idx, column=1, value=label)
        c2 = ws_kpi.cell(row=row_idx, column=2, value=val)
        c1.border = thin_border
        c2.border = thin_border
        if is_currency:
            c2.number_format = currency_format
            c2.font = bold_font
        if "Invalid" in label and val > 0:
            c1.fill = alert_fill
            c2.fill = alert_fill
        row_idx += 1

    # Top 5 Contractors Table
    row_idx += 2
    ws_kpi.cell(row=row_idx, column=1, value="Top 5 Contractors (by Netto Revenue)").font = Font(size=13, bold=True, color="1F4E78")
    row_idx += 1

    ws_kpi.cell(row=row_idx, column=1, value="Contractor").fill = header_fill
    ws_kpi.cell(row=row_idx, column=1).font = header_font
    ws_kpi.cell(row=row_idx, column=2, value="Revenue Netto").fill = header_fill
    ws_kpi.cell(row=row_idx, column=2).font = header_font
    row_idx += 1

    top_clients = df.groupby("contractor_name")["netto"].sum().sort_values(ascending=False).head(5)
    for client, amt in top_clients.items():
        c1 = ws_kpi.cell(row=row_idx, column=1, value=client)
        c2 = ws_kpi.cell(row=row_idx, column=2, value=amt)
        c1.border = thin_border
        c2.border = thin_border
        c2.number_format = currency_format
        row_idx += 1

    # ==========================================
    # SHEET 2: Cleaned Invoices
    # ==========================================
    ws_data = wb.create_sheet(title="Audited Invoices")
    ws_data.views.sheetView[0].showGridLines = True

    export_cols = [
        ("Invoice Number", "invoice_number"),
        ("Contractor", "contractor_name"),
        ("NIP", "nip_clean"),
        ("Issue Date", "standard_issue_date"),
        ("Netto", "netto"),
        ("VAT %", "vat_rate"),
        ("VAT Amount", "vat_amount"),
        ("Brutto", "brutto"),
        ("Audit Status", "audit_status"),
    ]

    # Header Row
    for col_idx, (col_label, _) in enumerate(export_cols, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data Rows
    for r_idx, row in df.iterrows():
        excel_row = r_idx + 2
        for c_idx, (_, col_field) in enumerate(export_cols, start=1):
            val = row[col_field]
            cell = ws_data.cell(row=excel_row, column=c_idx, value=val)
            cell.border = thin_border

            # Number formatting
            if col_field in ("netto", "vat_amount", "brutto"):
                cell.number_format = currency_format
            elif col_field == "vat_rate":
                cell.number_format = '0.0"%"'
            elif col_field == "audit_status":
                if val != "OK":
                    cell.fill = alert_fill
                    cell.font = Font(color="C00000", bold=True)
                else:
                    cell.font = Font(color="375623", bold=True)

    # Auto-fit column widths
    for sheet in [ws_kpi, ws_data]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
